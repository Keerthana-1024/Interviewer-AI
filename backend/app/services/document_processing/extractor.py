"""
Document processing: PDF (PyMuPDF) and DOCX (python-docx) extraction.
Validates file type and size. Never executes document content.
Treats document text as untrusted data.
"""
from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass, field
from typing import List, Optional

logger = logging.getLogger(__name__)

SUPPORTED_TYPES = {"application/pdf", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


@dataclass
class ExtractionResult:
    text: str
    pages: List[str] = field(default_factory=list)  # per-page text if available
    sections: dict = field(default_factory=dict)     # detected section → text
    file_type: str = ""
    word_count: int = 0
    success: bool = True
    error: Optional[str] = None

    def __post_init__(self):
        if self.text:
            self.word_count = len(self.text.split())


class DocumentProcessor:
    """
    Handles PDF and DOCX extraction.
    Raises ValueError for invalid files.
    Returns ExtractionResult — never silently returns empty text on error.
    """

    def __init__(self, max_bytes: int = 10 * 1024 * 1024):
        self.max_bytes = max_bytes

    def extract(self, file_bytes: bytes, filename: str) -> ExtractionResult:
        """
        Main entry point. Detects file type and dispatches.
        Validates size and type before processing.
        """
        self._validate_size(file_bytes, filename)

        ext = self._get_extension(filename)
        if ext not in SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file type: {ext!r}. Only PDF and DOCX are accepted."
            )

        if ext == ".pdf":
            return self._extract_pdf(file_bytes)
        else:
            return self._extract_docx(file_bytes)

    def _validate_size(self, file_bytes: bytes, filename: str):
        size = len(file_bytes)
        if size == 0:
            raise ValueError(f"File '{filename}' is empty.")
        if size > self.max_bytes:
            raise ValueError(
                f"File '{filename}' exceeds the maximum allowed size of "
                f"{self.max_bytes // (1024*1024)} MB."
            )

    @staticmethod
    def _get_extension(filename: str) -> str:
        filename = filename.lower()
        if filename.endswith(".pdf"):
            return ".pdf"
        if filename.endswith(".docx"):
            return ".docx"
        return "." + filename.rsplit(".", 1)[-1] if "." in filename else ""

    def _extract_pdf(self, file_bytes: bytes) -> ExtractionResult:
        try:
            import fitz  # PyMuPDF
        except ImportError:
            raise RuntimeError("PyMuPDF is not installed. Run: pip install pymupdf")

        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
        except Exception as e:
            return ExtractionResult(text="", success=False, error=f"Could not open PDF: {e}", file_type="pdf")

        if doc.page_count == 0:
            return ExtractionResult(text="", success=False, error="PDF has no pages.", file_type="pdf")

        pages: List[str] = []
        full_text_parts: List[str] = []

        for page_num in range(doc.page_count):
            try:
                page = doc[page_num]
                page_text = page.get_text("text")
                pages.append(page_text)
                full_text_parts.append(f"[Page {page_num + 1}]\n{page_text}")
            except Exception as e:
                logger.warning(f"Could not extract page {page_num + 1}: {e}")
                pages.append("")

        doc.close()

        full_text = "\n".join(full_text_parts)

        if not full_text.strip():
            return ExtractionResult(
                text="",
                success=False,
                error=(
                    "PDF text extraction produced empty content. "
                    "The file may be a scanned image PDF (OCR required). "
                    "Please provide a text-based PDF or DOCX."
                ),
                file_type="pdf",
            )

        sections = _detect_sections(full_text)
        return ExtractionResult(text=full_text, pages=pages, sections=sections, file_type="pdf")

    def _extract_docx(self, file_bytes: bytes) -> ExtractionResult:
        try:
            from docx import Document
        except ImportError:
            raise RuntimeError("python-docx is not installed. Run: pip install python-docx")

        try:
            doc = Document(io.BytesIO(file_bytes))
        except Exception as e:
            return ExtractionResult(text="", success=False, error=f"Could not open DOCX: {e}", file_type="docx")

        parts: List[str] = []
        for para in doc.paragraphs:
            if para.text.strip():
                # Include heading style info to help section detection
                style = para.style.name if para.style else ""
                if "Heading" in style:
                    parts.append(f"\n### {para.text.strip()} ###\n")
                else:
                    parts.append(para.text.strip())

        # Also extract tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(row_text)

        full_text = "\n".join(parts)

        if not full_text.strip():
            return ExtractionResult(text="", success=False, error="DOCX file is empty or has no readable text.", file_type="docx")

        sections = _detect_sections(full_text)
        return ExtractionResult(text=full_text, sections=sections, file_type="docx")


# ─── Section Detection ────────────────────────────────────────────────────────

SECTION_PATTERNS = {
    "skills": r"(?i)(skills?|technical\s+skills?|technologies|competencies|expertise|tools)",
    "experience": r"(?i)(experience|employment|work\s+history|professional\s+background|career)",
    "education": r"(?i)(education|academic|degree|university|college|school)",
    "projects": r"(?i)(projects?|portfolio|personal\s+projects?|academic\s+projects?)",
    "certifications": r"(?i)(certifications?|licenses?|credentials?|awards?)",
    "summary": r"(?i)(summary|objective|profile|about\s+me|overview)",
    "achievements": r"(?i)(achievements?|accomplishments?|highlights?|publications?)",
}


def _detect_sections(text: str) -> dict:
    """
    Heuristic section detection for resumes.
    Splits text around likely section headers.
    Returns {section_name: text_content}.
    """
    lines = text.splitlines()
    sections: dict = {}
    current_section = "header"
    current_lines: List[str] = []

    for line in lines:
        stripped = line.strip()
        matched_section = None

        # Check if this line looks like a section header
        if len(stripped) < 60 and stripped:
            for section_name, pattern in SECTION_PATTERNS.items():
                if re.search(pattern, stripped):
                    matched_section = section_name
                    break

        if matched_section:
            # Save previous section
            if current_lines:
                sections[current_section] = "\n".join(current_lines).strip()
            current_section = matched_section
            current_lines = []
        else:
            current_lines.append(line)

    # Save last section
    if current_lines:
        sections[current_section] = "\n".join(current_lines).strip()

    return sections
